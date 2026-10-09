"""Generates KaranPratap_ProjectReport.docx from results.json + figures/ (created by the notebook).
Usage: run the notebook first, then:  python build_report.py
"""
import json, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "KaranPratap_ProjectReport.docx"
R = json.load(open("results.json")) if os.path.exists("results.json") else None
PH = "[run notebook to fill]"


def g(fn, default=PH):
    if R is None:
        return default
    try:
        return fn(R)
    except Exception:
        return default


doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(1); s.top_margin = s.bottom_margin = Inches(0.9)


def h(text, lvl=1):
    doc.add_heading(text, level=lvl)


def p(text):
    para = doc.add_paragraph(text); para.paragraph_format.space_after = Pt(6); return para


def bullets(items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


def table(header, rows):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"
    for i, c in enumerate(header):
        t.rows[0].cells[i].text = c
        for r in t.rows[0].cells[i].paragraphs[0].runs: r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, c in enumerate(row): cells[i].text = str(c)
    doc.add_paragraph()


def fig(name, caption):
    path = f"figures/{name}.png"
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(5.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c = doc.add_paragraph(caption); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.runs[0].italic = True; c.runs[0].font.size = Pt(9)
    else:
        p(f"[Figure missing: {path} - run the notebook first]")


def money(x): return f"${x:,.0f}"


# ---------- Title ----------
t = doc.add_heading("Healthcare Analytics Project Report", 0); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
s = doc.add_paragraph("Data Cleaning, Exploratory Data Analysis and Billing Prediction"); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
a = doc.add_paragraph("Prepared by: Karan Pratap"); a.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

# ---------- Executive Summary ----------
h("1. Executive Summary")
clean = g(lambda r: r["cleaning"])
best = g(lambda r: r["best_regression_model"])
p("This project analyses a hospital admissions dataset to understand patient demographics, admission patterns and "
  "billing behaviour, and tests whether machine learning can predict billing amounts and test results from the "
  "available patient and admission attributes.")
bullets([
    f"Data quality: {g(lambda r: r['cleaning']['duplicates_removed'])} duplicate records and "
    f"{g(lambda r: r['cleaning']['negative_billing_removed'])} invalid (negative) billing records were removed, leaving "
    f"{g(lambda r: r['cleaning']['clean_rows'])} clean rows out of {g(lambda r: r['cleaning']['raw_rows'])}.",
    f"Billing: the average bill is {g(lambda r: money(r['eda']['mean_billing']))}; "
    f"{g(lambda r: r['top_billing_condition'][0])} has the highest average billing.",
    f"Modelling: the best regression model was {best} (R-squared {g(lambda r: [m['R2'] for m in r['regression'] if m['Model']==r['best_regression_model']][0])}). "
    f"Test-result classification accuracy was {g(lambda r: round(r['classification']['accuracy'], 3))} "
    f"against a chance level of {g(lambda r: round(r['classification']['chance_level'], 3))}.",
    "Recommendation: improve data quality controls at entry and enrich the data with clinical variables before "
    "relying on predictive models for billing or diagnosis support.",
])

# ---------- Dataset ----------
h("2. Dataset Details")
p("Source: Healthcare Dataset (Kaggle, prasad22/healthcare-dataset), file healthcare_dataset.csv. "
  "Each row represents one hospital admission. The dataset is synthetic, which is important when interpreting model results.")
table(["Column", "Description"], [
    ("Name, Age, Gender, Blood Type", "Patient demographics"),
    ("Medical Condition", "Primary condition (e.g. Cancer, Obesity, Diabetes, Asthma, Hypertension, Arthritis)"),
    ("Date of Admission / Discharge Date", "Admission and discharge dates"),
    ("Doctor, Hospital, Room Number", "Treating doctor, hospital and room"),
    ("Insurance Provider", "Patient's insurer"),
    ("Billing Amount", "Amount billed (target for regression)"),
    ("Admission Type", "Emergency, Urgent or Elective"),
    ("Medication", "Medication prescribed"),
    ("Test Results", "Normal, Abnormal or Inconclusive (target for classification)"),
])
p(f"Period covered: {g(lambda r: r['eda']['date_from'])} to {g(lambda r: r['eda']['date_to'])}. "
  f"Tools used: Python, pandas, NumPy, Matplotlib, Seaborn and scikit-learn.")

# ---------- Cleaning ----------
h("3. Data Cleaning")
table(["Step", "Action", "Impact"], [
    ("Column names", "Stripped spaces and replaced with underscores", "Consistent references"),
    ("Duplicates", "Dropped exact duplicate rows", f"{g(lambda r: r['cleaning']['duplicates_removed'])} rows removed"),
    ("Text formatting", "Standardised patient name capitalisation and trimmed whitespace", "Consistent names"),
    ("Dates", "Converted admission and discharge dates to datetime", "Enabled time features"),
    ("Invalid billing", "Removed records with negative billing amount", f"{g(lambda r: r['cleaning']['negative_billing_removed'])} rows removed"),
    ("Feature engineering", "Created Length_of_Stay, Admission_Month, Admission_Year, Age_Group", "New analysis features"),
])

# ---------- EDA ----------
h("4. Exploratory Data Analysis")
p(f"After cleaning, the average patient age is {g(lambda r: round(r['eda']['mean_age'], 1))} years, the average length of "
  f"stay is {g(lambda r: round(r['eda']['mean_los'], 1))} days and the average bill is {g(lambda r: money(r['eda']['mean_billing']))}. "
  f"The most common condition is {g(lambda r: r['eda']['most_common_condition'])} and the most common admission type is "
  f"{g(lambda r: r['eda']['most_common_admission'])}.")
fig("01_age_billing_distribution", "Figure 1: Age and billing amount distributions")
fig("02_condition_counts", "Figure 2: Patients by medical condition")
fig("03_avg_billing_condition", "Figure 3: Average billing by medical condition")
p(f"Average billing ranges from {g(lambda r: money(r['low_billing_condition'][1]))} ({g(lambda r: r['low_billing_condition'][0])}) "
  f"to {g(lambda r: money(r['top_billing_condition'][1]))} ({g(lambda r: r['top_billing_condition'][0])}).")
fig("04_admission_test_results", "Figure 4: Admission type and test results")
fig("05_billing_insurance", "Figure 5: Average billing by insurance provider")
fig("06_monthly_admissions", "Figure 6: Monthly admissions trend")
fig("07_los_agegroup", "Figure 7: Length of stay and billing by age group")
fig("08_correlation", "Figure 8: Correlation heatmap")
p(f"Correlation of billing with length of stay is {g(lambda r: round(r['billing_los_corr'], 3))} and with age is "
  f"{g(lambda r: round(r['billing_age_corr'], 3))}, i.e. how strongly billing moves with these variables.")

# ---------- Models ----------
h("5. Model Performance")
h("5.1 Regression: predicting billing amount", 2)
p("Linear Regression, Decision Tree and Random Forest were trained on an 80/20 train-test split using patient, "
  "condition, insurance, medication and stay-length features.")
if R:
    table(["Model", "MAE", "RMSE", "R-squared"],
          [(m["Model"], f"{m['MAE']:,.0f}", f"{m['RMSE']:,.0f}", m["R2"]) for m in R["regression"]])
    p(f"For reference, always predicting the training mean gives an MAE of {R['baseline_mae']:,.0f}.")
    best_r2 = max(m["R2"] for m in R["regression"])
    if best_r2 < 0.05:
        p("Interpretation: none of the models beat the naive mean baseline by a meaningful margin (R-squared close to zero). "
          "This indicates that the available features carry almost no information about the billed amount, which is consistent "
          "with a synthetic dataset where billing was generated independently of the other columns.")
    else:
        p(f"Interpretation: the best model explains about {best_r2:.0%} of the variance in billing, showing that the features carry "
          "a usable signal, although a substantial part of the variation remains unexplained.")
else:
    table(["Model", "MAE", "RMSE", "R-squared"], [(m, PH, PH, PH) for m in ("Linear Regression", "Decision Tree", "Random Forest")])
fig("09_actual_vs_predicted", "Figure 9: Actual vs predicted billing (best model)")
fig("10_feature_importance", "Figure 10: Top feature importances (Random Forest)")

h("5.2 Classification: predicting test results", 2)
p(f"A Random Forest classifier predicted Test Results with an accuracy of {g(lambda r: round(r['classification']['accuracy'], 3))}, "
  f"compared with a chance level of {g(lambda r: round(r['classification']['chance_level'], 3))} for three balanced classes.")
fig("11_confusion_matrix", "Figure 11: Confusion matrix for test results")

# ---------- Recommendations ----------
h("6. Recommendations")
bullets([
    "Data quality: add validation at data entry (no negative billing, duplicate detection, standard name formatting) so that cleaning "
    "is not repeated downstream.",
    "Billing transparency: review conditions and insurers with above-average billing and check whether the differences are explained "
    "by treatment intensity or by pricing inconsistencies.",
    "Capacity planning: use the monthly admissions trend and the admission-type mix to plan staffing and beds for emergency and urgent cases.",
    "Length of stay: track length of stay by condition and age group and flag unusually long stays for care-pathway review.",
    "Better predictive models: add clinical variables (diagnosis and procedure codes, lab values, comorbidities, prior admissions) "
    "before deploying any model for billing forecasts or test-result support.",
])

# ---------- Conclusion ----------
h("7. Conclusion and Limitations")
p("The project delivered a full analytics workflow: cleaning, exploratory analysis, regression and classification modelling, and "
  "actionable recommendations. The main limitation is the dataset itself: it is synthetic, so patterns and model scores should not be "
  "generalised to real hospitals. The workflow, however, can be reused directly on real hospital data.")

doc.save(OUT)
print("Saved", OUT)
